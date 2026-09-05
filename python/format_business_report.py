from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import DataBarRule
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from pathlib import Path


# ============================================================
# FILE PATH
# ============================================================

BASE_DIR = Path(r"D:\Clariant_Master_Data_Project")

FILE = (
    BASE_DIR
    / "excel"
    / "business_analysis_report.xlsx"
)


# ============================================================
# LOAD WORKBOOK
# ============================================================

wb = load_workbook(FILE)

print("=" * 70)
print("FORMATTING BUSINESS ANALYSIS REPORT")
print("=" * 70)


# ============================================================
# GENERAL STYLES
# ============================================================

header_fill = PatternFill(
    fill_type="solid",
    fgColor="1F4E78"
)

header_font = Font(
    bold=True,
    color="FFFFFF"
)

title_font = Font(
    bold=True,
    size=16
)

thin_border = Border(
    bottom=Side(style="thin")
)


# ============================================================
# FORMAT EACH SHEET
# ============================================================

for ws in wb.worksheets:

    # Header formatting
    for cell in ws[1]:

        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # Freeze first row
    ws.freeze_panes = "A2"

    # Auto filter
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions

    # Column widths
    for column_cells in ws.columns:

        max_length = 0
        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:

            if cell.value is not None:

                length = len(str(cell.value))

                if length > max_length:
                    max_length = length

        ws.column_dimensions[
            column_letter
        ].width = min(max_length + 2, 35)

    # Header row height
    ws.row_dimensions[1].height = 25


# ============================================================
# EXECUTIVE SUMMARY
# ============================================================

ws = wb["Executive Summary"]

ws["A3"] = "Master Data & Procurement Analytics"
ws["A3"].font = title_font

# Number formatting
if ws.max_row >= 2:

    ws["E2"].number_format = '#,##0.00'


# ============================================================
# SUPPLIER ANALYSIS
# ============================================================

ws = wb["Supplier Analysis"]

# Identify columns and apply number formatting
headers = {
    cell.value: cell.column
    for cell in ws[1]
}

if "total_spend" in headers:

    col = headers["total_spend"]

    for row in range(2, ws.max_row + 1):

        ws.cell(
            row=row,
            column=col
        ).number_format = '#,##0.00'

    # Data bar for supplier spend
    ws.conditional_formatting.add(
        f"{get_column_letter(col)}2:"
        f"{get_column_letter(col)}{ws.max_row}",
        DataBarRule(
            start_type="min",
            end_type="max",
            color="5B9BD5",
            showValue=True
        )
    )


# ============================================================
# MATERIAL ANALYSIS
# ============================================================

ws = wb["Material Analysis"]

headers = {
    cell.value: cell.column
    for cell in ws[1]
}

for field in [
    "price",
    "total_quantity",
    "total_spend"
]:

    if field in headers:

        col = headers[field]

        for row in range(2, ws.max_row + 1):

            ws.cell(
                row=row,
                column=col
            ).number_format = '#,##0.00'


# ============================================================
# CATEGORY ANALYSIS
# ============================================================

ws = wb["Category Analysis"]

headers = {
    cell.value: cell.column
    for cell in ws[1]
}

if "category_spend" in headers:

    col = headers["category_spend"]

    for row in range(2, ws.max_row + 1):

        ws.cell(
            row=row,
            column=col
        ).number_format = '#,##0.00'


# ============================================================
# CATEGORY SPEND CHART
# ============================================================

ws = wb["Category Analysis"]

chart = BarChart()

chart.type = "bar"

chart.title = "Procurement Spend by Category"

chart.y_axis.title = "Category"

chart.x_axis.title = "Procurement Spend"

headers = {
    cell.value: cell.column
    for cell in ws[1]
}

category_col = headers["category_name"]
spend_col = headers["category_spend"]

data = Reference(
    ws,
    min_col=spend_col,
    min_row=1,
    max_row=ws.max_row
)

categories = Reference(
    ws,
    min_col=category_col,
    min_row=2,
    max_row=ws.max_row
)

chart.add_data(
    data,
    titles_from_data=True
)

chart.set_categories(categories)

chart.height = 8
chart.width = 15

ws.add_chart(
    chart,
    "F2"
)


# ============================================================
# MONTHLY TREND
# ============================================================

ws = wb["Monthly Trend"]

headers = {
    cell.value: cell.column
    for cell in ws[1]
}

if "monthly_spend" in headers:

    col = headers["monthly_spend"]

    for row in range(2, ws.max_row + 1):

        ws.cell(
            row=row,
            column=col
        ).number_format = '#,##0.00'


# ============================================================
# MONTHLY PROCUREMENT LINE CHART
# ============================================================

ws = wb["Monthly Trend"]

chart = LineChart()

chart.title = "Monthly Procurement Spend"

chart.y_axis.title = "Procurement Spend"

chart.x_axis.title = "Month"

headers = {
    cell.value: cell.column
    for cell in ws[1]
}

month_col = headers["purchase_month"]
spend_col = headers["monthly_spend"]

data = Reference(
    ws,
    min_col=spend_col,
    min_row=1,
    max_row=ws.max_row
)

categories = Reference(
    ws,
    min_col=month_col,
    min_row=2,
    max_row=ws.max_row
)

chart.add_data(
    data,
    titles_from_data=True
)

chart.set_categories(categories)

chart.height = 8
chart.width = 15

ws.add_chart(
    chart,
    "F2"
)


# ============================================================
# DATA QUALITY
# ============================================================

ws = wb["Data Quality"]

for row in ws.iter_rows():

    for cell in row:

        if isinstance(cell.value, (int, float)):

            cell.number_format = '#,##0'


# ============================================================
# SAVE
# ============================================================

wb.save(FILE)


print()
print("=" * 70)
print("EXCEL REPORT FORMATTED SUCCESSFULLY")
print("=" * 70)
print()
print(f"File:")
print(FILE)
print()
print("Added:")
print(" - Professional headers")
print(" - Filters")
print(" - Frozen headers")
print(" - Number formatting")
print(" - Conditional formatting")
print(" - Category spend chart")
print(" - Monthly procurement trend chart")
print()
print("=" * 70)